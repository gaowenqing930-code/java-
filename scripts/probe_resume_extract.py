# -*- coding: utf-8 -*-
"""构造 docx/pdf 测试简历并直测 /interview/resume/extract"""
import io
import json
import time
import urllib.error
import urllib.request

from docx import Document

from reportlab.pdfgen import canvas as pdf_canvas

BASE = "http://localhost:8080/api"
TOKEN = None


def login():
    global TOKEN
    req = urllib.request.Request(BASE + "/auth/login", data=json.dumps({"username": "admin", "password": "admin123"}).encode(), method="POST")
    req.add_header("Content-Type", "application/json")
    with urllib.request.urlopen(req, timeout=10) as x:
        r = json.loads(x.read().decode())
    TOKEN = r["data"]["token"]


def extract(path, fname):
    import uuid
    boundary = "----b" + uuid.uuid4().hex
    with open(path, "rb") as f:
        body = f.read()
    parts = []
    parts.append(("--" + boundary).encode())
    parts.append(('Content-Disposition: form-data; name="file"; filename="%s"' % fname).encode())
    parts.append(b"Content-Type: application/octet-stream")
    parts.append(b"")
    parts.append(body)
    parts.append(("--" + boundary + "--").encode())
    payload = b"\r\n".join(parts)
    req = urllib.request.Request(BASE + "/interview/resume/extract", data=payload, method="POST")
    req.add_header("Content-Type", "multipart/form-data; boundary=" + boundary)
    req.add_header("Authorization", "Bearer " + TOKEN)
    t0 = time.time()
    try:
        with urllib.request.urlopen(req, timeout=60) as x:
            r = json.loads(x.read().decode())
        return r, round(time.time() - t0, 1)
    except urllib.error.HTTPError as e:
        return {"code": e.code, "msg": e.read().decode()[:200]}, round(time.time() - t0, 1)


login()

# --- docx ---
doc = Document()
doc.add_heading("张三的个人简历", level=1)
doc.add_paragraph("技术栈：Java、Spring Boot、MyBatis-Plus、MySQL、Redis、Vue3")
doc.add_paragraph("项目：AI智能校园平台助手，负责 JWT 鉴权与 Redis 缓存热点数据。")
p = "C:/Users/Administrator/AppData/Local/Temp/resume_test.docx"
doc.save(p)
r, t = extract(p, "resume_test.docx")
print("docx:", r.get("code"), "time:", t, "s")
print("  fileName:", r.get("data", {}).get("fileName"))
print("  text:", (r.get("data", {}).get("text") or "")[:120])

# --- pdf ---
q = "C:/Users/Administrator/AppData/Local/Temp/resume_test.pdf"
c = pdf_canvas.Canvas(q)
c.drawString(72, 760, "Zhang San Resume")
c.drawString(72, 740, "Tech Stack: Java, Spring Boot, MySQL, Redis")
c.drawString(72, 720, "Project: AI Campus Assistant")
c.save()
r, t = extract(q, "resume_test.pdf")
print("pdf:", r.get("code"), "time:", t, "s")
print("  text:", (r.get("data", {}).get("text") or "")[:120])

# --- 不支持格式 ---
w = "C:/Users/Administrator/AppData/Local/Temp/resume_test.txt"
open(w, "w", encoding="utf-8").write("hello")
r, t = extract(w, "resume_test.txt")
print("txt:", r.get("code"), "msg:", r.get("msg", "")[:60])
